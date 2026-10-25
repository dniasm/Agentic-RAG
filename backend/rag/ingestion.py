"""Ingestion pipeline: load -> clean -> chunk -> embed -> upsert into Qdrant."""
import hashlib
import logging
import re
import time
import uuid
from pathlib import Path

from langchain_core.documents import Document
from qdrant_client import models

from backend.config.settings import settings
from backend.rag import sparse, vectorstore

log = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


class IngestionError(Exception):
    pass


def file_sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _clean(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"-\n(\w)", r"\1", text)        # re-join hyphenated line breaks
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def load_file(path: Path) -> tuple[list[Document], int]:
    """Return one Document per page (PDF) or per file (others), plus the page count."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as e:
                raise IngestionError("PDF is password-protected") from e
        docs = []
        for i, page in enumerate(reader.pages, start=1):
            text = _clean(page.extract_text() or "")
            if text:
                docs.append(Document(page_content=text, metadata={"page": i}))
        return docs, len(reader.pages)

    if ext == ".docx":
        import docx

        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        text = _clean("\n".join(parts))
        return ([Document(page_content=text, metadata={"page": 1})] if text else []), 1

    if ext in {".txt", ".md"}:
        text = _clean(path.read_text(encoding="utf-8", errors="ignore"))
        return ([Document(page_content=text, metadata={"page": 1})] if text else []), 1

    raise IngestionError(f"Unsupported file type: {ext}")


def split_documents(docs: list[Document]) -> list[Document]:
    # Imported here, not at module level: the langchain_text_splitters package eagerly imports
    # sentence-transformers / transformers / torch when they're installed, which added 12-32 s to
    # every backend startup. Now that cost is paid on the first upload only (or not at all,
    # if the cross-encoder has already loaded torch in the background).
    from langchain_text_splitters import RecursiveCharacterTextSplitter

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=settings.chunk_size,
        chunk_overlap=settings.chunk_overlap,
        separators=["\n\n", "\n", ". ", " ", ""],
        add_start_index=True,
    )
    return [c for c in splitter.split_documents(docs) if len(c.page_content.strip()) > 20]


# Gemini's free tier counts EVERY text in a batch as one request, limited to 100 per minute. Sending a
# whole document at once therefore fails for any document over ~100 chunks, however long you wait.
EMBED_BATCH = 90
_RETRY_IN = re.compile(r"retry in ([\d.]+)s", re.I)


def embed_paced(texts: list[str], title: str) -> list[list[float]]:
    """Embed in batches under the per-minute limit; on a rate-limit error, wait as long as Google asks."""
    emb = vectorstore.get_embeddings()
    vectors: list[list[float]] = []
    for start in range(0, len(texts), EMBED_BATCH):
        batch = texts[start:start + EMBED_BATCH]
        for attempt in range(6):
            try:
                vectors += emb.embed_documents(batch, batch_size=EMBED_BATCH, task_type="RETRIEVAL_DOCUMENT",
                                               titles=[title] * len(batch))
                break
            except Exception as e:
                msg = str(e)
                if "RESOURCE_EXHAUSTED" not in msg or attempt == 5:
                    raise
                wait = float(_RETRY_IN.search(msg).group(1)) + 2 if _RETRY_IN.search(msg) else 60.0
                log.info("Embedding rate limit hit (chunk %d of %d); waiting %.0fs", start, len(texts), wait)
                time.sleep(wait)
    return vectors


def ingest_file(path: Path, *, user_id: str, doc_id: str, filename: str) -> dict:
    """Index a stored file for a user. Returns {"pages": int, "chunks": int}."""
    pages, page_count = load_file(path)
    if not pages:
        raise IngestionError("No extractable text found (the file may be scanned images).")

    chunks = split_documents(pages)
    texts = [c.page_content for c in chunks]
    vectors = embed_paced(texts, filename)

    points = []
    for idx, (chunk, vector) in enumerate(zip(chunks, vectors)):
        metadata = {
            "user_id": user_id,
            "doc_id": doc_id,
            "source": filename,
            "page": chunk.metadata.get("page"),
            "chunk_index": idx,
            "start_index": chunk.metadata.get("start_index"),
        }
        points.append(models.PointStruct(
            id=str(uuid.uuid5(uuid.NAMESPACE_URL, f"{doc_id}:{idx}")),
            vector={vectorstore.DENSE: vector, vectorstore.SPARSE: sparse.document_vector(chunk.page_content)},
            payload={"page_content": chunk.page_content, "metadata": metadata},
        ))

    vectorstore.upsert(points)
    return {"pages": page_count, "chunks": len(points)}
